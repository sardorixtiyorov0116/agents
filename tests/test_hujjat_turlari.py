"""Eski va «niqoblangan» formatlar: .xls, .doc (Word 97 / RTF / HTML), .docx jadvallari.

NEGA: menejer Telegramda «.xls ni o'qimaydi» degan javob oldi (2026-10-05).
Mijozlar TZ ni ko'pincha eski formatda yuboradi; «.doc» esa ichida uch xil
bo'ladi — Word 97, RTF, HTML («Word -> veb-sahifa»). Word hammasini ochadi.

Haqiqiy namunalar `tests/etalon_tz/hujjat_turlari/` da (git ga kirmaydi).
.doc jadvallari Word COM bilan olingan natija bilan solishtirilgan:
qatorlar soni va to'la kataklar 100% mos (vertikal birlashtirilgan katak
o'rnida biz bo'sh katak qoldiramiz — ustunlar siljimasin).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from kp.tz_jadval import jadval_oqi
from kp.word_doc import DocXatosi, hujjat_oqi, html_oqi, rtf_oqi

NAMUNA = Path(__file__).parent / "etalon_tz" / "hujjat_turlari"
namuna = pytest.mark.skipif(not NAMUNA.is_dir(), reason="namunalar yo'q (git ga kirmaydi)")


# --- sun'iy ---------------------------------------------------------------------


def _rtf_hex(matn: str) -> str:
    """Kirill -> RTF «\\'hh» (cp1251)."""
    return "".join(f"{chr(92)}'{b:02x}" for b in matn.encode("cp1251"))


def _rtf_unicode(matn: str) -> str:
    """Kirill -> RTF «\\uNNNN?»."""
    return "".join(f"{chr(92)}u{ord(c)}?" for c in matn)


def test_rtf_jadval_va_kirill():
    b = chr(92)
    rtf = (f"{{{b}rtf1{b}ansi{b}ansicpg1251{{{b}fonttbl{{{b}f0 Times;}}}}"
           f"{b}pard {_rtf_hex('Спецификация')}{b}par"
           f"{b}trowd{b}pard{b}intbl {_rtf_hex('Наименование')}{b}cell "
           f"{_rtf_hex('Кол-во')}{b}cell{b}row"
           f"{b}trowd{b}pard{b}intbl {_rtf_unicode('Решетка')} 300x150{b}cell 57{b}cell{b}row"
           f"{b}pard Tamom{b}par}}").encode("ascii")
    matn, jadvallar = rtf_oqi(rtf)
    assert jadvallar == [[["Наименование", "Кол-во"], ["Решетка 300x150", "57"]]]
    assert matn.startswith("Спецификация") and "Tamom" in matn


def test_html_doc_cp1251():
    html = ('<html><head><meta charset="windows-1251"><style>p{}</style></head><body>'
            '<p>Заявка</p><table><tr><td>Наименование</td><td>Кол-во</td></tr>'
            '<tr><td>Решетка&nbsp;АДН 300x150</td><td>57</td></tr></table></body></html>'
            ).encode("cp1251")
    matn, jadvallar = html_oqi(html)
    assert jadvallar == [[["Наименование", "Кол-во"], ["Решетка АДН 300x150", "57"]]]
    assert "Заявка" in matn


def test_docx_birlashtirilgan_katak_takrorlanmaydi(tmp_path):
    from docx import Document

    d = Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Наименование"
    t.cell(0, 2).text = "Кол-во"
    t.cell(1, 0).text = "Решетка"
    t.cell(1, 1).text = "300x150"
    t.cell(1, 2).text = "57"
    yol = tmp_path / "tz.docx"
    d.save(yol)
    _, jadvallar = hujjat_oqi(yol)
    assert jadvallar[0][0] == ["Наименование", "Кол-во"]
    assert jadvallar[0][1] == ["Решетка", "300x150", "57"]


def test_tur_mazmundan_aniqlanadi(tmp_path):
    """Kengaytmasi .doc, ichi HTML — Word shunday saqlaydi."""
    yol = tmp_path / "tz.doc"
    yol.write_bytes("<html><table><tr><td>Наименование</td><td>Кол-во</td></tr>"
                    "<tr><td>Клапан ДКС 300x200</td><td>59</td></tr></table></html>".encode())
    j = jadval_oqi(yol)
    assert [(q.nomi, q.miqdor) for q in j.qatorlar] == [("Клапан ДКС 300x200", 59)]


def test_tanilmagan_hujjat(tmp_path):
    yol = tmp_path / "x.doc"
    yol.write_bytes(b"oddiy matn")
    with pytest.raises(DocXatosi):
        hujjat_oqi(yol)


def test_excel_nom_davomi_alohida_mahsulot_emas(tmp_path):
    import openpyxl

    kitob = openpyxl.Workbook()
    v = kitob.active
    for q in [["№", "Наименование", "Тип", "Ед. изм.", "Кол-во"],
              [1, "Вентилятор радиальный с эл. двигателем", "ВЦ4-75-2,5", "комп.", 1],
              [None, "мощностью N=0,55 квт, п=2740 об/мин", "4АА6382", None, 1],
              [None, "Виброизоляторы", "ДО-40", None, 5]]:
        v.append(q)
    yol = tmp_path / "tz.xlsx"
    kitob.save(yol)
    q = jadval_oqi(yol).qatorlar
    assert len(q) == 2
    assert "мощностью N=0,55" in q[0].nomi
    assert (q[1].nomi, q[1].miqdor) == ("Виброизоляторы ДО-40", 5)


# --- haqiqiy namunalar ------------------------------------------------------------


@namuna
def test_xls_zayavka():
    from kp.tz import matn_ol
    from kp.tz_qoralama import fayldan_taklif

    j = jadval_oqi(NAMUNA / "zayavka_aeroport.xls")
    assert len(j.qatorlar) == 28            # nom davomlari birlashgan (33 -> 28)
    assert j.qatorlar[0].nomi.startswith("Вентилятор радиальный") and "N=0,55" in j.qatorlar[0].nomi
    assert "Аэропорт" in matn_ol(NAMUNA / "zayavka_aeroport.xls")
    assert fayldan_taklif(NAMUNA / "zayavka_aeroport.xls").javoblar["yol"] == "model"


@namuna
@pytest.mark.parametrize("fayl, qatorlar", [
    ("ol_klapanlar.doc", [18, 4]),          # Word 97, vertikal birlashtirilgan kataklar
    ("ol_reshetki.doc", [7, 4]),
    ("spek_2-6.doc", [72, 48]),             # aslida HTML
])
def test_doc_jadvallari_Word_bilan_bir_xil(fayl, qatorlar):
    _, jadvallar = hujjat_oqi(NAMUNA / fayl)
    assert [len(j) for j in jadvallar] == qatorlar


@namuna
def test_doc_birlashtirilgan_qator_aniq():
    """«КРВ: да / нет» — 4 va 2 katakli qatorlar; keyingi qatorlar surilmaydi."""
    _, jadvallar = hujjat_oqi(NAMUNA / "ol_reshetki.doc")
    t = jadvallar[0]
    assert t[3][:3] == ["3", "КРВ (клапан расхода воздуха)", "да"]
    assert [k for k in t[4] if k] == ["нет"]
    assert t[5][:2] == ["4", "RAL – цвет покраски по каталогу RAL"]


@namuna
def test_doc_matni_TZ_yoliga_tushadi():
    from kp.tz import matn_ol

    matn = matn_ol(NAMUNA / "ol_klapanlar.doc")
    assert "КПУ" in matn and "Количество, шт" in matn

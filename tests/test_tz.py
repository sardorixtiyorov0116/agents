"""TZ faylidan parametr ajratish — `kp/tz.py`.

BU YERDA MODEL CHAQIRILMAYDI. Testlar ikki narsani qo'riqlaydi:

  1. Fayldan MATN olish (sof kod) — PDF, DOCX, XLSX, TXT.
  2. Model qaytargan natijani shakl javoblariga aylantirish.

Eng muhim qoida: TZ dan olingan qiymat KP ga JIMGINA tushmaydi.
Har biri "topilganlar" ro'yxatiga yoziladi va menejerga ko'rsatiladi.
"""

from __future__ import annotations

import pytest

from kp.tz import (
    MAKS_HAJM,
    QOLLAB_QUVVATLANADI,
    ShaklTaklifi,
    TzNatija,
    TzXatosi,
    TzXona,
    matn_ol,
    shaklga_aylantir,
)

TURLAR = {"ofis", "ombor", "restoran", "sanuzel", "ishlab_chiqarish"}


# --- 1-bosqich: fayldan matn --------------------------------------------------


def test_txt_oqiladi(tmp_path):
    yol = tmp_path / "tz.txt"
    yol.write_text("Ombor 800 m2, balandligi 6 m", encoding="utf-8")

    assert "800" in matn_ol(yol)


def test_docx_JADVALI_ham_oqiladi(tmp_path):
    """TZ da parametrlar odatda JADVALDA turadi — ularsiz bo'lmaydi."""
    from docx import Document

    hujjat = Document()
    hujjat.add_paragraph("Texnik topshiriq")
    jadval = hujjat.add_table(rows=2, cols=3)
    jadval.rows[0].cells[0].text = "Xona"
    jadval.rows[0].cells[1].text = "Maydon"
    jadval.rows[1].cells[0].text = "Ombor"
    jadval.rows[1].cells[1].text = "800"
    yol = tmp_path / "tz.docx"
    hujjat.save(str(yol))

    matn = matn_ol(yol)

    assert "Texnik topshiriq" in matn
    assert "800" in matn, "jadval o'qilmadi"


def test_xlsx_oqiladi(tmp_path):
    import openpyxl

    kitob = openpyxl.Workbook()
    kitob.active.append(["Xona", "Maydon", "Balandlik"])
    kitob.active.append(["Sex", 1200, 8])
    yol = tmp_path / "tz.xlsx"
    kitob.save(str(yol))

    matn = matn_ol(yol)

    assert "1200" in matn


def test_notanish_format_RAD_etiladi(tmp_path):
    yol = tmp_path / "chizma.dwg"
    yol.write_bytes(b"\x00\x01")

    with pytest.raises(TzXatosi, match="formati o'qilmaydi"):
        matn_ol(yol)


def test_bosh_fayl_SABABINI_aytadi(tmp_path):
    """Skanerlangan PDF — rasm ichida matn. Jim qolmaymiz."""
    yol = tmp_path / "bosh.txt"
    yol.write_text("   \n\n  ", encoding="utf-8")

    with pytest.raises(TzXatosi, match="Skanerlangan|matn topilmadi"):
        matn_ol(yol)


def test_juda_katta_fayl_RAD_etiladi(tmp_path):
    yol = tmp_path / "katta.txt"
    yol.write_bytes(b"x" * (MAKS_HAJM + 1))

    with pytest.raises(TzXatosi, match="juda katta"):
        matn_ol(yol)


def test_yoq_fayl():
    with pytest.raises(TzXatosi, match="topilmadi"):
        matn_ol("yoq-bunday-fayl.pdf")


def test_qollab_quvvatlanadigan_formatlar():
    """DWG yo'q — chizmadan matn olib bo'lmaydi, buni va'da qilmaymiz."""
    assert QOLLAB_QUVVATLANADI == {".pdf", ".docx", ".xlsx", ".txt"}


# --- 2-bosqich: natijani shaklga aylantirish ----------------------------------


def test_obyekt_yoli_toldiriladi():
    natija = TzNatija(
        mijoz="TOY MCHJ", obyekt="Chilonzor SM", inn="123456789",
        xonalar=[TzXona(nomi="Zal", turi="restoran", maydon=500, balandlik=3.5,
                        odamlar=200)],
    )

    taklif = shaklga_aylantir(natija, TURLAR)

    assert taklif.javoblar["yol"] == "obyekt"
    assert taklif.javoblar["olcham"] == {"maydon": 500.0, "balandlik": 3.5}
    assert taklif.javoblar["xona_turi"] == "restoran"
    assert taklif.javoblar["odamlar"] == 200
    assert taklif.javoblar["inn"] == "123456789"


def test_topilgan_qiymatlar_MENEJERGA_KORSATILADI():
    """TZ dan olingan raqam KP ga JIMGINA tushmasligi kerak."""
    natija = TzNatija(
        mijoz="TOY MCHJ",
        xonalar=[TzXona(turi="ofis", maydon=250, balandlik=3)],
    )

    taklif = shaklga_aylantir(natija, TURLAR)

    matn = " ".join(taklif.topilganlar)
    assert "TOY MCHJ" in matn
    assert "250" in matn


def test_model_nomi_bolsa_MODEL_yoliga_otadi():
    """Model aniq bo'lsa hisoblash shart emas — narx darrov qidiriladi."""
    natija = TzNatija(mahsulotlar=["ВЦ 4-75-6,3-1", "РВН 300х300"])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert taklif.javoblar["yol"] == "model"
    assert [m["nomi"] for m in taklif.javoblar["mahsulotlar"]] == [
        "ВЦ 4-75-6,3-1", "РВН 300х300"]


def test_model_yolida_MIQDOR_ogohlantiriladi():
    """TZ da miqdor har doim ham aniq bo'lmaydi — 1 dona qo'yiladi."""
    natija = TzNatija(mahsulotlar=["ВЦ 4-75-6,3-1"])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert any("miqdor" in o for o in taklif.ogohlantirishlar)


# --- taxmin qilinmaydi --------------------------------------------------------


def test_balandlik_yoq_bolsa_TAXMIN_QILINMAYDI():
    """«Odatda 3 metr» deb qo'yish xato bo'lardi — hisob shunga tayanadi."""
    natija = TzNatija(xonalar=[TzXona(turi="ofis", maydon=250, balandlik=0)])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert "olcham" not in taklif.javoblar
    assert any("balandlig" in o for o in taklif.ogohlantirishlar)


def test_royxatda_YOQ_tur_OGOHLANTIRADI():
    """Noma'lum tur jim qabul qilinmaydi — normasi boshqacha bo'lishi mumkin."""
    natija = TzNatija(
        xonalar=[TzXona(nomi="Basseyn", turi="basseyn", maydon=250, balandlik=3)])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert taklif.javoblar["xonalar"][0]["turi"] == "ofis"   # zaxira norma
    assert any("basseyn" in o for o in taklif.ogohlantirishlar)


def test_notogri_INN_olinmaydi():
    """9 raqamdan boshqasi STIR emas."""
    natija = TzNatija(inn="12345", xonalar=[TzXona(maydon=100, balandlik=3)])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert "inn" not in taklif.javoblar


def test_maydonsiz_xona_OLINMAYDI():
    natija = TzNatija(xonalar=[TzXona(nomi="Zal", turi="ofis", balandlik=3)])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert "olcham" not in taklif.javoblar
    assert any("MAYDONI" in o for o in taklif.ogohlantirishlar)


# --- ko'p xonali TZ -----------------------------------------------------------


def test_HAMMA_xona_olinadi():
    """Ilgari faqat eng kattasi olinardi — qolgan xonalar yo'qolardi.

    3 xonali TZ dan bitta xonalik KP chiqardi va mijoz kam uskuna
    ko'rardi.
    """
    natija = TzNatija(xonalar=[
        TzXona(nomi="Ofis", turi="ofis", maydon=120, balandlik=3),
        TzXona(nomi="Sex", turi="ishlab_chiqarish", maydon=900, balandlik=6),
        TzXona(nomi="Ombor", turi="ombor", maydon=400, balandlik=6),
    ])

    taklif = shaklga_aylantir(natija, TURLAR)

    xonalar = taklif.javoblar["xonalar"]
    assert len(xonalar) == 3
    assert {x["nomi"] for x in xonalar} == {"Ofis", "Sex", "Ombor"}
    assert any("3 ta xona" in t for t in taklif.topilganlar)


def test_balandliksiz_xona_TASHLANADI_va_aytiladi():
    """Balandliksiz hisob qilib bo'lmaydi — lekin jim tashlanmaydi."""
    natija = TzNatija(xonalar=[
        TzXona(nomi="Ofis", turi="ofis", maydon=120, balandlik=3),
        TzXona(nomi="Podval", turi="ombor", maydon=400, balandlik=0),
    ])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert len(taklif.javoblar["xonalar"]) == 1
    assert any("Podval" in o and "OLINMADI" in o
               for o in taklif.ogohlantirishlar)


def test_TZ_dagi_havo_sarfi_SOLISHTIRISHGA_qoyiladi():
    """TZ da tayyor sarf bo'lsa — u loyihachi hisobi, e'tiborsiz qolmasin."""
    natija = TzNatija(xonalar=[
        TzXona(turi="ofis", maydon=250, balandlik=3, havo_sarfi=3200)])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert any("3200" in o for o in taklif.ogohlantirishlar)


def test_bosh_natija_bosh_taklif():
    taklif = shaklga_aylantir(TzNatija(), TURLAR)

    assert not taklif.bormi
    assert isinstance(taklif, ShaklTaklifi)


def test_mahsulot_topilsa_XONALAR_jim_tashlanmaydi():
    """JONLI HOLAT (2026-08-26): «Albom №6 — Ustaxona binosi» chizmasi.

    Loyiha chizmasida ODATDA ikkalasi ham bo'ladi: xonalar
    ekspliikatsiyasi va loyihachi tanlagan uskuna. Ilgari mahsulot
    topilishi bilan 12 ta xona jimgina yo'qolardi va menejer ular
    o'qilganini umuman bilmasdi.
    """
    natija = TzNatija(
        mahsulotlar=["VK 200", "VKP 4E"],
        xonalar=[
            TzXona(nomi="Рабочая комната", turi="ofis", maydon=258, balandlik=3),
            TzXona(nomi="Сварочный цех", turi="ishlab_chiqarish", maydon=21, balandlik=3),
        ],
    )

    taklif = shaklga_aylantir(natija, TURLAR)

    assert taklif.javoblar["yol"] == "model"      # uskuna ustun
    assert any("2 ta xona ham bor" in o for o in taklif.ogohlantirishlar)
    assert any("Obyektni tasvirlayman" in o for o in taklif.ogohlantirishlar)


def test_xonasiz_mahsulotda_ortiqcha_xabar_YOQ():
    natija = TzNatija(mahsulotlar=["VK 200"])

    taklif = shaklga_aylantir(natija, TURLAR)

    assert not any("xona ham bor" in o for o in taklif.ogohlantirishlar)


# --- Aniqlanmagan xona turi ---------------------------------------------------
#
# JONLI XATO (2026-08-28): zaxira model (`gemini-3.5-flash-lite`)
# «Торговый зал» uchun turni BO'SH qaytardi va tizim jimgina `ofis`
# qo'ydi. Ogohlantirish faqat NOTO'G'RI tur uchun yozilardi — bo'sh tur
# hech qanday iz qoldirmasdi.
#
# Tur havo almashinuvi normasini belgilaydi: restoran (8.0 karrali)
# o'rniga ofis (3.0) olinsa sarf 2,7 BAROBAR kam chiqadi.


def _taklif(turi: str):
    from hisob import normalar
    from kp.tz import TzNatija, TzXona, shaklga_aylantir

    natija = TzNatija(mijoz="Sinov", xonalar=[
        TzXona(nomi="Zal", turi=turi, maydon=850.0, balandlik=4.0, odamlar=0)])
    return shaklga_aylantir(natija, set(normalar()))


def test_BOSH_tur_ogohlantirish_beradi():
    taklif = _taklif("")

    assert any("ANIQLANMADI" in o or "aniqlanmadi" in o
               for o in taklif.ogohlantirishlar)


def test_NOTOGRI_tur_ham_ogohlantirish_beradi():
    taklif = _taklif("yoq_bunday_tur")

    assert any("ro'yxatda yo'q" in o for o in taklif.ogohlantirishlar)


def test_TOGRI_tur_ogohlantirmaydi():
    taklif = _taklif("restoran")

    assert taklif.ogohlantirishlar == []


def test_aniqlanmagan_tur_ROYXATDA_korinadi():
    """Menejer qaysi norma ishlatilganini ko'rmasa, xatoni topa olmaydi."""
    taklif = _taklif("")

    zal = [t for t in taklif.topilganlar if t.startswith("Zal")]
    assert zal and "ofis" in zal[0]


def test_aniqlangan_tur_ham_ROYXATDA_korinadi():
    taklif = _taklif("restoran")

    zal = [t for t in taklif.topilganlar if t.startswith("Zal")]
    assert zal and "(restoran)" in zal[0]


def test_bosh_turda_ofis_normasi_ishlatiladi():
    """Oqim to'xtamasin — lekin buni AYTAMIZ (yuqoridagi testlar)."""
    taklif = _taklif("")

    assert taklif.javoblar["xonalar"][0]["turi"] == "ofis"

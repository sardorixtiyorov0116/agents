"""Tender summasi O'Z VALYUTASIDA va texnik topshiriqni yuklash.

JONLI HOLAT (2026-09-14): etender da 866 ochiq lotdan 54 tasi DOLLARDA,
10 tasi YEVRODA. Valyuta o'qilmasdi va hamma summa "so'm" deb chiqardi:
81 205,70 dollarlik lot "81 206 so'm" bo'lib ko'rinardi.

Tarmoqqa CHIQILMAYDI — soxta transport bilan. Modul `oz_transporti`
bilan belgilangan: conftest dagi tarmoq qo'riqchisi aynan shu
funksiyalarni almashtiradi, bu yerda esa HAQIQIYSI sinaladi.
"""

from __future__ import annotations

import json

import httpx
import pytest

import integrations.valyuta as valyuta
from integrations.brend import begona_brendlar
from integrations.tender_manba import (
    EtenderManba,
    TenderXatosi,
    _hujjatni_oqi,
    etender_lot_hujjatlari,
    etender_lot_id,
    mcuz_lot_hujjatlari,
    mcuz_lot_id,
)

from .test_tender_manba import ETENDER_YOZUV, _soxta_etender

pytestmark = pytest.mark.oz_transporti


@pytest.fixture(autouse=True)
def kesh_tozala():
    valyuta._KESH_HAMMASI = None
    yield
    valyuta._KESH_HAMMASI = None


# --- pul matni ---------------------------------------------------------------


@pytest.mark.parametrize("summa,kod,kutilgan", [
    (12_500_000, "UZS", "12 500 000 so'm"),
    (1234.56, "UZS", "1 235 so'm"),            # so'm avvalgidek yumaloqlanadi
    (81205.7, "USD", "81 205,70 dollar"),       # sent yo'qolmaydi
    (700000.0, "EUR", "700 000 yevro"),
    (5, "GBP", "5 GBP"),                        # nomi yo'q valyuta — ISO kod
    (None, "USD", "—"),
])
def test_pul_matni(summa, kod, kutilgan):
    assert valyuta.pul_matni(summa, kod) == kutilgan


def test_somga():
    kurslar = {"UZS": 1.0, "USD": 11765.76}

    assert valyuta.somga(100, "USD", kurslar) == pytest.approx(1_176_576)
    assert valyuta.somga(5, "UZS", None) == 5
    # Kursi noma'lum — taxmin qilinmaydi.
    assert valyuta.somga(100, "GBP", kurslar) is None
    assert valyuta.somga(None, "UZS", kurslar) is None


# --- Markaziy bank kurslari ----------------------------------------------------


class _Javob:
    def __init__(self, tana, kod=200):
        self._tana = tana
        self.status_code = kod

    def raise_for_status(self):
        if self.status_code >= 400:
            raise httpx.HTTPStatusError("xato", request=None, response=None)

    def json(self):
        return self._tana


class _Mijoz:
    def __init__(self, javob=None, xato=None):
        self.javob = javob
        self.xato = xato

    async def get(self, url, timeout=None):
        if self.xato:
            raise self.xato
        return self.javob


async def test_kurslar_nominal_bilan_oqiladi():
    mijoz = _Mijoz(_Javob([
        {"Ccy": "USD", "Nominal": "1", "Rate": "11765.76"},
        {"Ccy": "EUR", "Nominal": "1", "Rate": "13638.87"},
        {"Ccy": "JPY", "Nominal": "10", "Rate": "800"},   # 10 birlik uchun
        {"Ccy": "XXX", "Rate": "buzuq"},                  # tashlanadi
    ]))

    kurslar = await valyuta.markaziy_bank_kurslari(mijoz=mijoz, keshdan=False)

    assert kurslar["UZS"] == 1.0
    assert kurslar["USD"] == 11765.76
    assert kurslar["JPY"] == 80.0
    assert "XXX" not in kurslar


async def test_kurslar_bosh_javobda_xato():
    with pytest.raises(valyuta.KursXatosi):
        await valyuta.markaziy_bank_kurslari(mijoz=_Mijoz(_Javob([])), keshdan=False)


async def test_kurslar_tarmoq_xatosi_yutilmaydi():
    mijoz = _Mijoz(xato=httpx.ConnectError("ulanmadi"))
    with pytest.raises(valyuta.KursXatosi):
        await valyuta.markaziy_bank_kurslari(mijoz=mijoz, keshdan=False)


# --- etender: lot valyutasi ---------------------------------------------------


async def test_etender_DOLLAR_lot_dollar_bolib_qoladi():
    yozuv = {**ETENDER_YOZUV, "cost": 81205.7, "currency_id": 14,
             "currency_codeabc": "USD", "currency_name": "Доллар США"}

    elonlar = await EtenderManba(mijoz=_soxta_etender([yozuv])).elonlar(chek=10)

    assert elonlar[0].summa == 81205.7
    assert elonlar[0].valyuta == "USD"
    assert elonlar[0].qisqa()["valyuta"] == "USD"


async def test_etender_valyutasiz_lot_SOM():
    elonlar = await EtenderManba(mijoz=_soxta_etender([ETENDER_YOZUV])).elonlar(chek=10)

    assert elonlar[0].valyuta == "UZS"


# --- etender: texnik topshiriq ------------------------------------------------


def test_lot_id_faqat_etender_havolasidan():
    assert etender_lot_id("https://etender.uzex.uz/lot/510351") == "510351"
    assert etender_lot_id("https://tender.mc.uz/tender-list/tender/5/view") is None
    assert etender_lot_id("") is None


def _pdf(matn: str) -> bytes:
    """Bir betli, matnli, HAQIQIY PDF — pypdf o'qiy oladigan."""
    oqim = f"BT /F1 12 Tf 72 720 Td ({matn}) Tj ET".encode("latin-1")
    obyektlar = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(oqim)).encode() + b" >>\nstream\n" + oqim + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    chiqish = bytearray(b"%PDF-1.4\n")
    ofsetlar = []
    for tartib, obyekt in enumerate(obyektlar, 1):
        ofsetlar.append(len(chiqish))
        chiqish += f"{tartib} 0 obj\n".encode() + obyekt + b"\nendobj\n"
    xref = len(chiqish)
    chiqish += f"xref\n0 {len(obyektlar) + 1}\n0000000000 65535 f \n".encode()
    for ofset in ofsetlar:
        chiqish += f"{ofset:010d} 00000 n \n".encode()
    chiqish += (
        f"trailer\n<< /Size {len(obyektlar) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(chiqish)


def _docx(*qatorlar: str, jadval: list[list[str]] | None = None) -> bytes:
    import io

    from docx import Document

    hujjat = Document()
    for qator in qatorlar:
        hujjat.add_paragraph(qator)
    if jadval:
        t = hujjat.add_table(rows=len(jadval), cols=len(jadval[0]))
        for i, kataklar in enumerate(jadval):
            for j, katak in enumerate(kataklar):
                t.cell(i, j).text = katak
    oqim = io.BytesIO()
    hujjat.save(oqim)
    return oqim.getvalue()


TZ_YOLI = "/files/2026/8/20/tz.pdf"
SHARTNOMA_YOLI = "files/2026/9/3/shartnoma.docx"

LOT = {
    "name": "Konditsionerlash VRF tizimlarining tashqi bloklariga joriy ta'mirlash",
    "addon_description": "Texnik xizmat ko'rsatish",
    "budget_products": json.dumps([
        {"Product_Name": "Установка кондиционера",
         "Description": "Первичная диагностика наружного блока VRF"},
    ], ensure_ascii=False),
    "valuation_name": "Ball usuli",
    "tech_coef": 60.0,
    "cost_coef": 40.0,
    "advance_payment_perc": 15.0,
    "payment_type_name": "Oldindan to'lov",
    "pledge_value": 1.0,          # birligi noma'lum — shartlarga KIRMAYDI
    "js_fields": [{
        "label": "O'xshash tovar tajribasi",
        "description": "Tajribasi yetarli bo'lmasa ishtirokchi chetlashtiriladi",
    }],
    "tech_file_path": TZ_YOLI,
    "tech_doc_file_path": SHARTNOMA_YOLI,
    "tech_doc_file_name": "shartnoma.docx",
}


def _lot_mijozi(lot=LOT, fayllar=None, yuklanganlar=None):
    """`fayllar`: yo'l -> baytlar yoki HTTP kodi."""
    if fayllar is None:
        fayllar = {
            TZ_YOLI: _pdf("Marka/model: ARV6-H610/SR1MV"),
            SHARTNOMA_YOLI: _docx(
                "Ehtiyot qismlar Ijrochi hisobidan",
                jadval=[["205 dona", "420 000,00"]],
            ),
        }

    def ishlov(soro: httpx.Request) -> httpx.Response:
        if soro.method == "GET" and soro.url.path == "/api/common/GetTrade/510351/0":
            return httpx.Response(200, json=lot)
        if soro.method == "POST" and soro.url.path == "/api/common/DownloadFile":
            yol = soro.url.params.get("path")
            if yuklanganlar is not None:
                yuklanganlar.append(yol)
            fayl = fayllar.get(yol, 404)
            if isinstance(fayl, int):
                return httpx.Response(fayl)
            return httpx.Response(200, content=fayl)
        return httpx.Response(404)

    return httpx.AsyncClient(transport=httpx.MockTransport(ishlov))


async def test_lot_hujjatlari_kartochka_shartlar_va_fayllar():
    yuklanganlar: list[str] = []

    hujjat = await etender_lot_hujjatlari(
        "510351", mijoz=_lot_mijozi(yuklanganlar=yuklanganlar)
    )

    # Yo'l API bergan ko'rinishda uzatiladi; DOCX ham yuklanadi.
    assert yuklanganlar == [TZ_YOLI, SHARTNOMA_YOLI]
    assert [nomi for nomi, _ in hujjat.fayllar] == ["tz.pdf", "shartnoma.docx"]
    assert hujjat.oqilmagan == []
    # Kartochkadan: mahsulot tavsifi va TO'XTATUVCHI baholash mezoni.
    assert "Первичная диагностика наружного блока VRF" in hujjat.karta
    assert "chetlashtiriladi" in hujjat.karta
    # Shartlar KODDA ko'chiriladi; zakalat birligi noma'lum — yo'q.
    assert hujjat.shartlar == [
        "Baholash: Ball usuli (texnika 60%, narx 40%)",
        "Avans: 15%",
        "To'lov: Oldindan to'lov",
    ]
    # Shartnoma loyihasidagi shart va jadval ham matnda.
    assert "Ehtiyot qismlar Ijrochi hisobidan" in hujjat.matn
    assert "205 dona | 420 000,00" in hujjat.matn
    assert begona_brendlar(hujjat.matn)[0].korinish == "AUX (ARV6-H610/SR1MV)"


async def test_lot_topilmasa_XATO():
    def ishlov(soro):
        return httpx.Response(404)

    with pytest.raises(TenderXatosi):
        await etender_lot_hujjatlari(
            "510351", mijoz=httpx.AsyncClient(transport=httpx.MockTransport(ishlov))
        )


async def test_fayl_orniga_sahifa_kelsa_oqilmagan_deb_yoziladi():
    """Sayt fayl o'rniga SPA sahifasini qaytaradi.

    Bu "hujjatda hech narsa yo'q" deb jimgina o'tib ketmasin — sababi
    bilan `oqilmagan` ga tushadi.
    """
    sahifa = b"<!doctype html><html></html>"

    hujjat = await etender_lot_hujjatlari(
        "510351", mijoz=_lot_mijozi(fayllar={TZ_YOLI: sahifa, SHARTNOMA_YOLI: sahifa})
    )

    assert hujjat.fayllar == []
    assert len(hujjat.oqilmagan) == 2
    assert all("sayt sahifasi" in sabab for sabab in hujjat.oqilmagan)


async def test_bitta_fayl_yuklanmasa_qolgani_oqiladi():
    hujjat = await etender_lot_hujjatlari(
        "510351",
        mijoz=_lot_mijozi(fayllar={TZ_YOLI: _pdf("Chiller 50 kW"), SHARTNOMA_YOLI: 500}),
    )

    assert [nomi for nomi, _ in hujjat.fayllar] == ["tz.pdf"]
    assert hujjat.oqilmagan == ["shartnoma.docx: yuklanmadi (500)"]


def test_ZIP_ichidagi_hujjatlar_oqiladi(monkeypatch):
    """511261 TZ si ZIP da, 511606 niki ZIP ichidagi papkada.

    Windows arxivatori UTF-8 bayrog'ini qo'ymaydi va kirillni CP866 da
    yozadi — nom buzilmasligi kerak.
    """
    import io
    import zipfile

    ichki = io.BytesIO()
    with zipfile.ZipFile(ichki, "w") as arxiv:
        arxiv.writestr("a.pdf", _pdf("ichki"))

    monkeypatch.setattr(
        zipfile.ZipInfo, "_encodeFilenameFlags",
        lambda self: (self.filename.encode("cp437"), self.flag_bits & ~0x800),
    )
    oqim = io.BytesIO()
    with zipfile.ZipFile(oqim, "w") as arxiv:
        kirill = "ТЗ/Техническое задание.pdf".encode("cp866").decode("cp437")
        arxiv.writestr(kirill, _pdf("Chiller 50 kW"))
        arxiv.writestr("shartnoma.docx", _docx("Avans 15%"))
        arxiv.writestr("eski.doc", b"\xd0\xcf\x11\xe0" + b"\0" * 64)
        arxiv.writestr("ichki.zip", ichki.getvalue())

    oqilgan, oqilmagan = _hujjatni_oqi("tz.zip", oqim.getvalue())

    assert [nomi for nomi, _ in oqilgan] == ["Техническое задание.pdf", "shartnoma.docx"]
    assert "Chiller 50 kW" in oqilgan[0][1]
    assert "Avans 15%" in oqilgan[1][1]
    assert oqilmagan == [
        "eski.doc: eski Word/Excel formati (.doc, .xls) o'qilmaydi",
        "ichki.zip: arxiv ichidagi arxiv ochilmaydi",
    ]


def test_matnsiz_skaner_PDF_ochiq_aytiladi():
    oqilgan, oqilmagan = _hujjatni_oqi("skan.pdf", _pdf(""))

    assert oqilgan == []
    assert oqilmagan == ["skan.pdf: matn qatlami yo'q (skanerlangan rasm)"]


# --- tender.mc.uz: lot hujjatlari ---------------------------------------------


def test_mcuz_lot_id():
    assert mcuz_lot_id("https://tender.mc.uz/tender-list/tender/287840/view") == "287840"
    assert mcuz_lot_id("https://etender.uzex.uz/lot/510351") is None


async def test_mcuz_LOYIHA_xizmati_kartochkadan_korinadi():
    """JONLI HOLAT: 287840 "joriy ta'mirlash" deb chiqdi, aslida loyiha-smeta."""
    lot = {
        "name": "Oliy sud oshxonasi ventilyatsiyasini joriy ta'mirlash",
        "object_type": {"id": 4, "name": "Проектно-изыскательный"},
        "service_type": {"id": 5, "name": "Текущий ремонт"},
        "plan": {"items": [
            {"enkt_code": {"name": "Услуга по разработке проектно-сметных работ"}},
        ]},
        "specializations": [
            {"name": "Разработка проектно-сметной документации", "toifa_required": 1},
        ],
        "end_term_work_days": 10,
        "customer": {"name": "Buyurtmachi"},              # fayl emas
        "loyiha_pdf": {"file": "/storage/1/loyiha.pdf", "file_name": "Илова 1,2.pdf"},
    }

    def ishlov(soro: httpx.Request) -> httpx.Response:
        if soro.url.path == "/api/tenders/287840":
            return httpx.Response(200, json={"result": {"data": lot}})
        if soro.url.path == "/storage/1/loyiha.pdf":
            return httpx.Response(200, content=_pdf("DEFEKTNIY AKT: ventilyator 1 sht"))
        return httpx.Response(404)

    hujjat = await mcuz_lot_hujjatlari(
        "287840", mijoz=httpx.AsyncClient(transport=httpx.MockTransport(ishlov))
    )

    assert hujjat.shartlar == [
        "Obyekt turi: Проектно-изыскательный",
        "Ish turi: Текущий ремонт",
        "Xarid predmeti (ENKT): Услуга по разработке проектно-сметных работ",
        "Talab qilinadigan ixtisoslik: Разработка проектно-сметной документации"
        " — toifa talab qilinadi",
        "Bajarish muddati: 10 kun",
    ]
    assert [nomi for nomi, _ in hujjat.fayllar] == ["Илова 1,2.pdf"]
    assert "ventilyator 1 sht" in hujjat.matn

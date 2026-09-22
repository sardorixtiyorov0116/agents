"""Tijorat menejeri Temur — KP, narx to'qimaslik, raqam takrorlanmasligi."""

from __future__ import annotations

import asyncio
import json
from datetime import date, timedelta

import pytest

from app.agentlar.tijorat_menejeri import TijoratMenejeri
from app.baza import Baza
from app.kontraktlar import kontraktlarni_yukla
from app.konvert import Holat, Ishonch
from kp import KP, Mijoz, Qator, Shartlar
from kp.narx import NarxRoyxati
from presenter.agentlar import agent_matni

from .soxta import SoxtaLlm, javob, matn_bloki, soxta_api, soxta_bilim

KATALOG = [
    {"id": 1, "name_uz": "Kanal ventilyatori ВК-250П", "category": {"name_uz": "Kanal"},
     "models": [{"name": "ВК-250П"}], "quantity": 10},
]


MENEJER = {"ism": "Яхшибоев Бауржон", "telefon": "+998 90 099 12 60"}


@pytest.fixture(autouse=True)
def menejer_malum(monkeypatch):
    """Standart holda menejer aniq — testlar KP mazmuniga qaratilsin.

    Menejer NOMA'LUM bo'lgan yo'l alohida sinaladi (`test_menejer_*`).
    """
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "menejer_uchun", lambda tg_id: dict(MENEJER))


@pytest.fixture
def kontrakt():
    return kontraktlarni_yukla()["proposal-builder"]


@pytest.fixture
async def baza(tmp_path):
    b = Baza(tmp_path / "temur.db")
    await b.tayyorla()
    return b


def temur_javobi(**ustama):
    malumot = {
        "mijoz": {"nomi": '"Oqtepa Qurilish" MChJ', "aloqa": "+998901234567", "manzil": ""},
        "til": "uz",
        "qatorlar": [
            {
                "nomi": "Kanal ventilyatori ВК-250П",
                "spetsifikatsiya": "230 V, 210 Vt",
                "miqdor": 5,
                "birlik": "dona",
                "izoh": "",
            }
        ],
        "topilmagan_mahsulotlar": [],
        "taklif_qilingan": [],
        "yetkazish": "",
        "maxsus_shartlar": [],
        "sorash_kerak": [],
    }
    malumot.update(ustama)
    return SoxtaLlm([javob([matn_bloki(json.dumps(malumot, ensure_ascii=False))])])


def temur_yasa(kontrakt, baza, llm, katalog=True):
    return TijoratMenejeri(
        kontrakt=kontrakt,
        llm=llm,
        baza=baza,
        api=soxta_api(qidir_keng=KATALOG) if katalog else soxta_api(qidir_keng=[]),
        qidiruv_manbasi=soxta_bilim(("KP shabloni", "Standart KP matni …")),
    )


# --- narx manbai -------------------------------------------------------------


def narx_royxati(narx=2_400_000, sana=None, kun=90):
    """sana=None -> bugungi; sana="" -> sanasiz (eskirgan hisoblanadi)."""
    return NarxRoyxati({
        "valyuta": "UZS",
        "qqs_foizi": 12,
        "amal_qilish_kuni": kun,
        "narxlar": [{
            "kod": "ВК-250П",
            "nomi": "Kanal ventilyatori ВК-250П",
            "narx": narx,
            "sana": date.today().isoformat() if sana is None else sana,
        }],
        "shartlar": {"tolov": "100% oldindan"},
    })


def test_narx_kirill_lotin_boyicha_topiladi():
    r = narx_royxati()
    assert r.top("VK-250P") is not None
    assert r.top("ВК-250П").narx == 2_400_000
    assert r.top("Noma'lum mahsulot") is None


def test_eskirgan_narx_ishlatilmaydi():
    eski = (date.today() - timedelta(days=200)).isoformat()
    narx = narx_royxati(sana=eski)
    topilgan = narx.top("ВК-250П")

    assert topilgan.eskirgan is True
    assert topilgan.bormi is False, "eskirgan narx ishlatilmasligi kerak"


def test_sanasiz_narx_eskirgan_hisoblanadi():
    narx = narx_royxati(sana="")
    assert narx.top("ВК-250П").bormi is False


def test_nol_narx_narx_yoq_degani():
    """0 — 'narx kiritilmagan', 'bepul' emas."""
    assert narx_royxati(narx=0).top("ВК-250П").bormi is False


# --- KP modeli ---------------------------------------------------------------


def kp_yasa(qatorlar, **kw):
    asos = dict(
        raqam="KP-2026-0001",
        sana=date(2026, 7, 29),
        mijoz=Mijoz(nomi="Test"),
        qatorlar=qatorlar,
        shartlar=Shartlar(),
        rekvizitlar={"nomi": "Test MChJ"},
        valyuta="UZS",
        qqs_foizi=12,
    )
    asos.update(kw)
    return KP(**asos)


# --- hujjat tartibi (namunaviy KP bilan bir xil bo'lishi shart) --------------


def _hujjat_tartibi(yol):
    """DOCX tanasidagi elementlar tartibi: ('P', matn) yoki ('T', qator soni)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    hujjat = docx.Document(str(yol))
    tartib = []
    for bola in hujjat.element.body.iterchildren():
        belgi = bola.tag.split("}")[-1]
        if belgi == "p":
            matn = Paragraph(bola, hujjat).text.strip()
            if matn:
                tartib.append(("P", matn))
        elif belgi == "tbl":
            tartib.append(("T", len(Table(bola, hujjat).rows)))
    return tartib


def test_hujjat_tartibi_namunadagidek(tmp_path):
    """Menejer bloki imzodan KEYIN turishi kerak (namunaviy КП_12951-26 dagidek).

    Ilgari u jadval bilan shartlar orasiga tushib qolardi.
    """
    from kp.hujjat import docx_yasa

    kp = kp_yasa(
        [Qator(nomi="ПВН 500-250-2", miqdor=5, birlik_narx=1_700_000, qqs_foizi=12)],
        rekvizitlar={
            "nomi": "Test MChJ",
            "direktor": "Расулов Ж.С.",
            "direktor_lavozimi": "Директор",
            "menejer": "Яхшибоев Бауржон",
            "menejer_telefon": "+998 90 099 12 60",
        },
        shartlar_matni=["Shartlar matni."],
        ogohlantirishlar=["sinov ogohlantirishi"],
    )
    tartib = _hujjat_tartibi(docx_yasa(kp, tmp_path / "tartib.docx"))
    yorliqlar = [t[1] if t[0] == "P" else "JADVAL" for t in tartib]

    def joyi(bolak):
        return next(i for i, y in enumerate(yorliqlar) if bolak in str(y))

    # Standart til — rus (namunaviy hujjat shu tilda).
    assert kp.til == "ru"

    # 1) sarlavha -> 2) mahsulot jadvali -> 3) shartlar -> 4) imzo -> 5) menejer
    assert joyi("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ") < joyi("Shartlar matni")
    assert joyi("Shartlar matni") < joyi("С уважением")
    assert joyi("С уважением") < joyi("Яхшибоев"), "menejer imzodan KEYIN"
    assert joyi("Яхшибоев") < joyi("+998 90 099 12 60")
    # Menejer bloki — hujjatning eng oxiri
    assert joyi("+998 90 099 12 60") == len(yorliqlar) - 1


def test_abzaslar_jadval_bilan_bir_chiziqda(tmp_path):
    """"С уважением," jadvaldan ichkariga surilib qolmasligi kerak.

    Sabab: `SimpleDocTemplate` freymga 6pt ichki bo'shliq qo'shadi, jadval
    esa undan chetda qolardi — natijada 2 mm farq paydo bo'lardi.
    """
    pymupdf = pytest.importorskip("pymupdf")

    from kp.hujjat import CHEKKA_MM, pdf_yasa

    kp = kp_yasa(
        [Qator(nomi="A", miqdor=1, birlik_narx=1000)],
        rekvizitlar={
            "nomi": "Test MChJ",
            "direktor": "Расулов Ж.С.",
            "direktor_lavozimi": "Директор",
            "menejer": "Яхшибоев Бауржон",
            "menejer_telefon": "+998 90 099 12 60",
        },
    )
    hujjat = pymupdf.open(str(pdf_yasa(kp, tmp_path / "abzas.pdf")))

    chapdagilar: dict[str, float] = {}
    for blok in hujjat[0].get_text("dict")["blocks"]:
        for qator in blok.get("lines", []):
            matn = "".join(b["text"] for b in qator["spans"]).strip()
            for kalit in ("уважением", "Директор", "Менеджер", "тел:"):
                if kalit in matn:
                    chapdagilar.setdefault(kalit, qator["bbox"][0] / 72 * 25.4)

    assert len(chapdagilar) == 4, f"topilmadi: {chapdagilar}"
    for kalit, mm_qiymat in chapdagilar.items():
        assert abs(mm_qiymat - CHEKKA_MM) < 0.5, f"{kalit} — {mm_qiymat:.1f}mm"


# --- hujjat tili -------------------------------------------------------------

KATALOG_IKKI_TILDA = [
    {
        "id": 31,
        "name_uz": "Suv bilan ishlaydigan kanalli isitgich ПВН",
        "name_ru": "Канальный водяной нагреватель ПВН",
        "category": {"name_uz": "To'rtburchak kanallar", "name_ru": "Прямоугольные каналы"},
        "models": [{"id": 28, "name": "ПВН 500-250-2", "price": "1700000"}],
        "quantity": 4,
    }
]


def test_til_sorovdan_aniqlanadi():
    from app.agentlar.tijorat_menejeri import til_aniqla

    assert til_aniqla("KP ni ruscha qilib ber") == "ru"
    assert til_aniqla("KP ni o'zbekcha qilib ber") == "uz"
    assert til_aniqla("KP rus tilida bo'lsin") == "ru"
    assert til_aniqla("hujjat o'zbek tilida kerak") == "uz"
    assert til_aniqla("сделай на русском") == "ru"
    assert til_aniqla("uzbekcha KP") == "uz"
    # Aytilmagan — tizim o'zi hal qiladi
    assert til_aniqla("Olmaliq AGMK uchun KP tayyorla") is None


def test_ozbekiston_soz_til_deb_qabul_qilinmaydi():
    """"O'zbekiston" — mijoz nomi, til ko'rsatmasi emas."""
    from app.agentlar.tijorat_menejeri import til_aniqla

    assert til_aniqla("O'zbekiston Temir Yo'llari uchun KP") is None
    assert til_aniqla("Узбекистон Темир Йуллари учун КП") is None


def test_oxirgi_til_korsatmasi_kuchda():
    from app.agentlar.tijorat_menejeri import til_aniqla

    assert til_aniqla("ruscha emas, o'zbekcha qilib ber") == "uz"


@pytest.mark.asyncio
async def test_ruscha_sorovda_katalog_ruscha_beriladi(
    kontrakt, baza, monkeypatch, tmp_path
):
    """Ruscha KP da mahsulot nomi ham backenddan RUSCHA olinadi."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi()
    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=llm, baza=baza,
        api=soxta_api(qidir_keng=KATALOG_IKKI_TILDA),
        qidiruv_manbasi=soxta_bilim(),
    )
    await temur.ishla("ПВН 500-250-2 uchun KP ni ruscha qilib ber")

    xabar = llm.chaqiruvlar[0]["messages"][0]["content"]
    assert "Канальный водяной нагреватель" in xabar
    assert "Suv bilan ishlaydigan" not in xabar
    assert "HUJJAT TILI: ru" in xabar


@pytest.mark.asyncio
async def test_ozbekcha_sorovda_katalog_ozbekcha_beriladi(
    kontrakt, baza, monkeypatch, tmp_path
):
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi()
    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=llm, baza=baza,
        api=soxta_api(qidir_keng=KATALOG_IKKI_TILDA),
        qidiruv_manbasi=soxta_bilim(),
    )
    await temur.ishla("ПВН 500-250-2 uchun KP ni o'zbekcha qilib ber")

    xabar = llm.chaqiruvlar[0]["messages"][0]["content"]
    assert "Suv bilan ishlaydigan kanalli isitgich" in xabar
    assert "Канальный водяной нагреватель" not in xabar
    assert "HUJJAT TILI: uz" in xabar


@pytest.mark.asyncio
async def test_sorovdagi_til_model_javobidan_ustun(
    kontrakt, baza, monkeypatch, tmp_path
):
    """Model "ru" desa ham, menejer o'zbekcha so'ragan bo'lsa — o'zbekcha."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    k = await temur_yasa(kontrakt, baza, temur_javobi(til="ru")).ishla(
        "KP ni o'zbekcha qilib ber"
    )
    assert k.natija["til"] == "uz"


def test_hujjat_matnlari_ikki_tilda(tmp_path):
    """O'zbekcha KP da sarlavha, ustunlar va kompaniya nomi ham o'zbekcha."""
    from kp.hujjat import _jadval_malumoti, _kompaniya, _m

    uchun = dict(
        rekvizitlar={"nomi": "ООО «CLIMAVENT»", "nomi_uz": "«CLIMAVENT» MChJ"},
        shartlar_matni=["Shartlar"],
    )
    ru = kp_yasa([Qator(nomi="A", miqdor=1, birlik_narx=1000)], til="ru", **uchun)
    uz = kp_yasa([Qator(nomi="A", miqdor=1, birlik_narx=1000)], til="uz", **uchun)

    assert _m(ru, "bosh") == "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ"
    assert _m(uz, "bosh") == "TIJORAT TAKLIFI"
    assert _m(ru, "murojaat") == "Руководителю"
    assert _m(uz, "murojaat") == "Rahbariga"
    assert _jadval_malumoti(ru)[0][0][1] == "Наименование"
    assert _jadval_malumoti(uz)[0][0][1] == "Nomi"
    assert _jadval_malumoti(ru)[0][-1][0] == "Итого:"
    assert _jadval_malumoti(uz)[0][-1][0] == "Jami:"
    # Kompaniya nomi ham tilga qarab
    assert _kompaniya(ru) == "ООО «CLIMAVENT»"
    assert _kompaniya(uz) == "«CLIMAVENT» MChJ"


def test_rekvizitlar_tilga_qarab_tanlanadi():
    """Imzo bloki ham tarjima qilinadi (`_uz` qo'shimchasi)."""
    from kp.hujjat import _rekvizit

    uchun = dict(
        rekvizitlar={
            "direktor": "Расулов Ж.С.", "direktor_uz": "Rasulov J.S.",
            "direktor_lavozimi": "Директор", "direktor_lavozimi_uz": "Direktor",
        }
    )
    ru = kp_yasa([Qator(nomi="A", miqdor=1)], til="ru", **uchun)
    uz = kp_yasa([Qator(nomi="A", miqdor=1)], til="uz", **uchun)

    assert _rekvizit(ru, "direktor_lavozimi") == "Директор"
    assert _rekvizit(uz, "direktor_lavozimi") == "Direktor"
    assert _rekvizit(uz, "direktor") == "Rasulov J.S."


def test_haqiqiy_rekvizitlarda_ikkala_til_bor():
    """Sozlama fayli to'ldirilgan bo'lsin — aks holda uz KP yarim ruscha chiqadi."""
    from kp import rekvizitlar

    r = rekvizitlar()
    for kalit in ("nomi", "direktor", "direktor_lavozimi"):
        assert r.get(kalit), f"{kalit} to'ldirilmagan"
        assert r.get(f"{kalit}_uz"), f"{kalit}_uz to'ldirilmagan"


def test_ozbekcha_nomi_yoq_bolsa_ruschasi_ishlatiladi():
    from kp.hujjat import _kompaniya

    kp = kp_yasa([Qator(nomi="A", miqdor=1)], til="uz",
                 rekvizitlar={"nomi": "ООО «CLIMAVENT»"})
    assert _kompaniya(kp) == "ООО «CLIMAVENT»"


# --- menejer bloki -----------------------------------------------------------


def menejer_sozla(monkeypatch, tmp_path, royxatdagi=None):
    """Menejer manbalarini nolga tushiradi (yoki oldindan kiritilganini beradi)."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul, "menejerlar", lambda: {})
    monkeypatch.setattr(modul, "menejer_uchun", lambda tg_id: royxatdagi)
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    yozilgan: dict = {}
    monkeypatch.setattr(
        modul, "hujjatlarni_yasa", lambda kp, papka: yozilgan.update(kp=kp) or {}
    )
    return yozilgan


@pytest.mark.asyncio
async def test_menejer_nomalum_bolsa_soraydi(kontrakt, baza, monkeypatch, tmp_path):
    """KP mijozga ketadi — noto'g'ri odamning nomi bilan chiqmasin."""
    from app.agentlar.tijorat_menejeri import MENEJER_SAVOLI

    menejer_sozla(monkeypatch, tmp_path)

    k = await temur_yasa(kontrakt, baza, temur_javobi()).ishla(
        "KP tayyorla, ruscha", {"telegram_id": 111}
    )

    assert k.holat is Holat.ANIQLIK_KERAK
    assert k.natija["savollar"] == [MENEJER_SAVOLI]
    # Bot javobni qaysi savolga bog'lashini shu belgidan biladi
    assert k.natija["kerak"] == "menejer"
    assert "telefon" in MENEJER_SAVOLI.lower()


@pytest.mark.asyncio
async def test_saqlangan_menejer_qayta_soralmaydi(kontrakt, baza, monkeypatch, tmp_path):
    """Javob bir marta beriladi — keyingi KP larda bazadan olinadi."""
    yozilgan = menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "Aziz Karimov", "+998 90 111 11 11")

    k = await temur_yasa(kontrakt, baza, temur_javobi()).ishla(
        "KP tayyorla, ruscha", {"telegram_id": 111}
    )

    assert k.holat is Holat.TUGADI
    # KP ruscha - ism ham kirillda (transliteratsiya: kp/harf.py)
    assert yozilgan["kp"].rekvizitlar["menejer"] == "Азиз Каримов"
    assert yozilgan["kp"].rekvizitlar["menejer_telefon"] == "+998 90 111 11 11"


@pytest.mark.asyncio
async def test_menejer_har_xodimga_alohida(kontrakt, baza, monkeypatch, tmp_path):
    """Bot bir necha kishida ishlaydi — javoblar aralashib ketmasin."""
    yozilgan = menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "Aziz Karimov", "+998 90 111 11 11")
    await baza.menejer_yoz(222, "Bekzod Aliyev", "+998 90 222 22 22")

    await temur_yasa(kontrakt, baza, temur_javobi()).ishla("KP ruscha", {"telegram_id": 222})
    assert yozilgan["kp"].rekvizitlar["menejer"] == "Бекзод Алиев"

    await temur_yasa(kontrakt, baza, temur_javobi()).ishla("KP ruscha", {"telegram_id": 111})
    assert yozilgan["kp"].rekvizitlar["menejer"] == "Азиз Каримов"


@pytest.mark.asyncio
async def test_oldindan_kiritilgan_xodimdan_soralmaydi(
    kontrakt, baza, monkeypatch, tmp_path
):
    """`rekvizitlar.yaml` da yozilgan bo'lsa savol berilmaydi."""
    yozilgan = menejer_sozla(
        monkeypatch, tmp_path,
        royxatdagi={"ism": "Yaml Xodim", "telefon": "+998 90 333 33 33"},
    )

    k = await temur_yasa(kontrakt, baza, temur_javobi()).ishla("KP ruscha", {"telegram_id": 111})

    assert k.holat is Holat.TUGADI
    assert yozilgan["kp"].rekvizitlar["menejer"] == "Ямл Ходим"


@pytest.mark.asyncio
async def test_sorovda_aytilgan_menejer_ustun(kontrakt, baza, monkeypatch, tmp_path):
    yozilgan = menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "Aziz Karimov", "+998 90 111 11 11")

    llm = temur_javobi(menejer="Dilshod Rahimov")
    await temur_yasa(kontrakt, baza, llm).ishla("KP ruscha", {"telegram_id": 111})

    assert yozilgan["kp"].rekvizitlar["menejer"] == "Дилшод Раҳимов"


def test_telefon_ajratish():
    """Foydalanuvchi javobidan ism va raqam ajratiladi."""
    from bot.asosiy import _menejer_ajrat

    assert _menejer_ajrat("Яхшибоев Бауржон +998 90 099 12 60") == (
        "Яхшибоев Бауржон", "+998 90 099 12 60",
    )
    assert _menejer_ajrat("+998901234567 Aziz Karimov") == (
        "Aziz Karimov", "+998901234567",
    )
    assert _menejer_ajrat("Aziz Karimov, +998 (90) 123-45-67")[0] == "Aziz Karimov"
    # Raqamsiz javob qabul qilinmaydi — savol kuchda qoladi
    assert _menejer_ajrat("Aziz Karimov") == ("", "")
    assert _menejer_ajrat("") == ("", "")


def test_menejerlar_royxati_yamldan_oqiladi():
    """Haqiqiy sozlama fayli to'g'ri o'qilishi (tuzilma buzilmasin)."""
    from kp import menejerlar, rekvizitlar

    royxat = menejerlar()
    assert isinstance(royxat, dict)
    for kalit, qiymat in royxat.items():
        assert kalit.isdigit(), f"kalit Telegram ID bo'lishi kerak: {kalit}"
        assert qiymat["ism"]
    # `menejerlar` matnli rekvizitlarga aralashib ketmasin
    assert "menejerlar" not in rekvizitlar()



def test_uzun_spetsifikatsiya_qisqartiriladi():
    """Uzun tavsif jadvalni buzadi va hujjatni ikkinchi sahifaga chiqaradi."""
    from kp.hujjat import SPETS_MAKS, _jadval_malumoti, _qisqa_spetsifikatsiya

    uzun = (
        "Suv bilan ishlaydigan kanalli isitgich. Kanal o'lchami: 500x250 mm. "
        "Qatorlar soni: 2. Havo bo'yicha unumdorlik: 600-1600 m3/soat. "
        "Issiqlik quvvati: 9,7-26,1 kVt. Suv sarfi: 0,35-0,93 m3/soat. "
        "Gidravlik qarshilik: 0,82-2,89 kPa. Korpus: rux qoplamali po'lat."
    )
    qisqa = _qisqa_spetsifikatsiya(uzun)

    assert len(qisqa) <= SPETS_MAKS + 1
    assert qisqa.endswith((".", "…")), "yarim jumla qolmasligi kerak"
    assert uzun.startswith(qisqa.rstrip(".…")[:40])

    # Jadvalda ham qisqargan holda chiqadi
    kp = kp_yasa([
        Qator(nomi="ПВН 500-250-2", spetsifikatsiya=uzun, miqdor=1, birlik_narx=1000)
    ])
    qatorlar, _ = _jadval_malumoti(kp)
    assert len(qatorlar[1][1]) < len(uzun)


def test_qisqa_spetsifikatsiya_tegilmaydi():
    from kp.hujjat import _qisqa_spetsifikatsiya

    qisqa = "500x250 mm, 2 qator, 9,7-26,1 kVt"
    assert _qisqa_spetsifikatsiya(qisqa) == qisqa


def test_blanka_toliq_kenglikda(tmp_path):
    """Logotip kichraytirilmasin: telefon/email/manzil RASM ICHIDA turadi."""
    import re

    from pypdf import PdfReader

    from kp.hujjat import BLANKA_KENGLIGI_MM, pdf_yasa

    kp = kp_yasa(
        [Qator(nomi="A", miqdor=1, birlik_narx=1000)],
        rekvizitlar={"nomi": "Test MChJ", "logo": "shablonlar/logo.png"},
    )
    yol = pdf_yasa(kp, tmp_path / "blanka.pdf")

    oqim = PdfReader(str(yol)).pages[0].get_contents().get_data().decode("latin-1")
    kengliklar = [
        float(m.group(1))
        for m in re.finditer(r"([\d.]+) [\d.\-]+ [\d.\-]+ ([\d.]+) [\d.\-]+ [\d.\-]+ cm", oqim)
        if float(m.group(1)) > 50 and float(m.group(2)) > 10
    ]
    assert kengliklar, "blanka rasmi topilmadi"
    # Namunada 193 mm; 1 mm farq yaxlitlashdan.
    assert abs(kengliklar[0] / 72 * 25.4 - BLANKA_KENGLIGI_MM) < 2


def test_qalin_shrift_ishlaydi():
    """`<b>` tegi shrift OILASI ro'yxatdan o'tmasa jimgina ishlamay qoladi."""
    from reportlab.lib.fonts import tt2ps

    from kp.hujjat import _shriftlar

    oddiy, qalin = _shriftlar()
    if oddiy == qalin:  # tizimda qalin TTF yo'q — tekshiradigan narsa qolmadi
        return
    # `<b>` aynan shu moslashtirish orqali qalin shriftga o'tadi.
    assert tt2ps(oddiy, 1, 0) == qalin


def test_jami_qatorida_qqs_bosh(tmp_path):
    """Namunada yakuniy qatorda QQS katagi bo'sh — faqat summa va jami."""
    from kp.hujjat import _jadval_malumoti

    kp = kp_yasa([Qator(nomi="A", miqdor=1, birlik_narx=1_000_000, qqs_foizi=12)])
    qatorlar, _ = _jadval_malumoti(kp)
    oxirgi = qatorlar[-1]

    assert oxirgi[0] == "Итого:"
    assert oxirgi[5] == kp.son(kp.summa)
    assert oxirgi[6] == "", "QQS katagi namunadagidek bo'sh qolishi kerak"
    assert oxirgi[7] == kp.son(kp.jami)


def test_narxsiz_qator_summaga_kirmaydi():
    kp = kp_yasa([
        Qator(nomi="A", miqdor=2, birlik_narx=1000),
        Qator(nomi="B", miqdor=3),  # narx yo'q
    ])

    assert kp.summa == 2000
    assert kp.qqs == 240
    assert kp.jami == 2240
    assert kp.toliq_narxmi is False
    assert len(kp.narxsiz_qatorlar) == 1


def test_hech_qaysi_qatorda_narx_yoq():
    kp = kp_yasa([Qator(nomi="A", miqdor=1)])

    assert kp.summa is None
    assert kp.jami is None
    assert kp.pul(kp.jami) == "—"


def test_pul_formati():
    kp = kp_yasa([Qator(nomi="A", miqdor=1, birlik_narx=12_500_000)])
    assert kp.pul(kp.summa) == "12 500 000 UZS"


def test_jadval_soni_namunaviy_formatda():
    """Namunaviy KP dagidek: 9 372 000,00"""
    assert KP.son(9_372_000) == "9 372 000,00"
    assert KP.son(1_124_640) == "1 124 640,00"
    assert KP.son(None) == ""


def test_qator_boyicha_qqs():
    """Namunaviy KP da har qatorda QQS va QQSli summa alohida ustun."""
    q = Qator(nomi="A", miqdor=1, birlik_narx=9_372_000, qqs_foizi=12)

    assert q.jami == 9_372_000
    assert q.qqs == pytest.approx(1_124_640)
    assert q.qqs_bilan == pytest.approx(10_496_640)


def test_narxsiz_qatorda_qqs_ham_bosh():
    q = Qator(nomi="A", miqdor=2, qqs_foizi=12)
    assert (q.jami, q.qqs, q.qqs_bilan) == (None, None, None)


def test_namunaviy_kp_summasi():
    """Haqiqiy KP (№ 12951/6) bilan bir xil summa chiqishi kerak."""
    kp = kp_yasa([
        Qator(nomi="ВЦ 4-75", miqdor=1, birlik_narx=9_372_000, qqs_foizi=12),
        Qator(nomi="ВО 30-160-6,3", miqdor=1, birlik_narx=10_850_000, qqs_foizi=12),
        Qator(nomi="ВО 30-160-8", miqdor=1, birlik_narx=20_570_000, qqs_foizi=12),
    ])

    assert KP.son(kp.summa) == "40 792 000,00"
    assert KP.son(kp.jami) == "45 687 040,00"


def test_jadvalda_itogo_yorligi_yoqolmaydi():
    """Birlashtirilgan katakda faqat BIRINCHI katak matni ko'rinadi."""
    from kp.hujjat import _jadval_malumoti

    kp = kp_yasa([Qator(nomi="A", miqdor=1, birlik_narx=1000, qqs_foizi=12)], til="ru")
    qatorlar, _ = _jadval_malumoti(kp)

    assert qatorlar[-1][0] == "Итого:", "yorliq birinchi katakda bo'lishi shart"
    assert len(qatorlar[0]) == 8, "namunaviy KP da 8 ustun"


def test_raqam_shakli():
    """Namunaviy KP raqami: 12951/6"""
    shakl = "{tartib}/{oy}"
    assert shakl.format(tartib=12951, oy=6, yil=2026, qisqa_yil=26) == "12951/6"


# --- raqamlash ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_kp_raqami_takrorlanmaydi(baza):
    raqamlar = [await baza.kp_raqam_ol() for _ in range(5)]

    assert len(set(raqamlar)) == 5, "raqamlar takrorlanmasligi kerak"
    assert raqamlar[0].startswith("KP-")
    # Ketma-ket o'sadi
    tartiblar = [int(r.rsplit("-", 1)[1]) for r in raqamlar]
    assert tartiblar == sorted(tartiblar)


@pytest.mark.asyncio
async def test_bir_vaqtda_sorasa_ham_takrorlanmaydi(baza):
    """Parallel so'rovlar bir xil raqam ololmaydi (UNIQUE indeks)."""
    raqamlar = await asyncio.gather(*[baza.kp_raqam_ol() for _ in range(10)])
    assert len(set(raqamlar)) == 10


# --- agent -------------------------------------------------------------------


@pytest.mark.asyncio
async def test_kp_tayyorlanadi_va_tasdiq_kerak(kontrakt, baza, monkeypatch, tmp_path):
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi()
    k = await temur_yasa(kontrakt, baza, llm).ishla("5 dona ВК-250П uchun KP, ruscha")

    assert k.kim == "proposal-builder"
    assert k.holat is Holat.TUGADI
    # O'zgarmas qoida: KP har doim tasdiqdan o'tadi
    assert k.tasdiq_kerak is True
    assert "MIJOZGA YUBORILMADI" in k.izoh
    assert k.natija["raqam"].startswith("KP-")
    assert k.natija["jami"] == pytest.approx(5 * 2_400_000 * 1.12)


@pytest.mark.asyncio
async def test_narx_backenddan_olinadi(kontrakt, baza, monkeypatch, tmp_path):
    """Birlamchi manba — backend: sayt va KP bir xil narxni ko'rsatishi kerak."""
    import app.agentlar.tijorat_menejeri as modul

    # YAML da boshqa narx turibdi — backend uni yengishi kerak.
    monkeypatch.setattr(modul, "narxlar", lambda: narx_royxati(narx=999_000))
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=temur_javobi(), baza=baza,
        # Narx to'liq katalogdan alohida so'raladi — qidiruv natijasidan emas.
        api=soxta_api(qidir_keng=KATALOG, model_narxi=(3_100_000, "ВК-250П")),
        qidiruv_manbasi=soxta_bilim(("KP shabloni", "matn")),
    )
    k = await temur.ishla("5 dona ВК-250П uchun KP, ruscha")

    # BACKEND NARXI QQS BILAN keladi (prays ustuni «Цена USD с НДС»),
    # KP jadvali esa QQSsiz summadan tuziladi — shuning uchun qatorda
    # QQS ajratilgan narx turadi.
    assert k.natija["qatorlar"][0]["birlik_narx"] == pytest.approx(3_100_000 / 1.12)
    # ENG MUHIMI: KP dagi JAMI backend narxiga TENG bo'lsin — mijoz
    # saytda va KP da bir xil raqam ko'rishi kerak. Ilgari bu yerda
    # 12% ortiqcha chiqardi.
    assert k.natija["jami"] == pytest.approx(5 * 3_100_000)
    # Backenddan olingan bo'lsa "yaml dan olindi" ogohlantirishi chiqmaydi.
    assert not any("narxlar.yaml" in o for o in k.natija["ogohlantirishlar"])


@pytest.mark.asyncio
async def test_backendda_narx_yoq_bolsa_yaml_zaxira(kontrakt, baza, monkeypatch, tmp_path):
    """Backendda price=0 — narxlar.yaml zaxira sifatida ishlaydi va aytiladi."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    k = await temur_yasa(kontrakt, baza, temur_javobi()).ishla("5 dona ВК-250П uchun KP, ruscha")

    assert k.natija["qatorlar"][0]["birlik_narx"] == 2_400_000
    assert any("narxlar.yaml" in o for o in k.natija["ogohlantirishlar"])


@pytest.mark.asyncio
async def test_narx_topilmasa_bosh_qoladi(kontrakt, baza, monkeypatch, tmp_path):
    """O'zgarmas qoida: narx to'qib chiqarilmaydi."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", lambda: narx_royxati(narx=0))
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi()
    k = await temur_yasa(kontrakt, baza, llm).ishla("5 dona ВК-250П uchun KP, ruscha")

    assert k.natija["qatorlar"][0]["birlik_narx"] is None
    assert k.natija["summa"] is None
    assert k.natija["toliq_narxmi"] is False
    assert k.ishonch is Ishonch.ORTA
    # "Topilmadi" emas, KIMDAN so'rash kerakligi yozilsin: narx
    # buxgalteriyada turadi (standartniki SAP'da, nostandartniki qo'lda).
    assert any(
        "buxgalteriyadan so'ralsin" in o for o in k.natija["ogohlantirishlar"]
    )


@pytest.mark.asyncio
async def test_model_narx_qaytarsa_ham_etiborga_olinmaydi(kontrakt):
    """Narx maydoni model sxemasida umuman yo'q — to'qib bo'lmaydi."""
    from app.agentlar.tijorat_menejeri import TemurQator

    maydonlar = set(TemurQator.model_fields)
    assert "birlik_narx" not in maydonlar
    assert not any("narx" in m for m in maydonlar)


@pytest.mark.asyncio
async def test_mahsulot_topilmasa_aniq_aytadi(kontrakt, baza, monkeypatch, tmp_path):
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi(
        topilmagan_mahsulotlar=["Jetfun-9000"],
        taklif_qilingan=["Kanal ventilyatori ВК-250П"],
    )
    k = await temur_yasa(kontrakt, baza, llm, katalog=False).ishla("Jetfun-9000 uchun KP, ruscha")

    assert "Jetfun-9000" in k.natija["topilmagan_mahsulotlar"]
    assert k.natija["taklif_qilingan"]
    assert any("topilmadi" in o for o in k.natija["ogohlantirishlar"])


@pytest.mark.asyncio
async def test_mahsulotsiz_sorovda_aniqlashtirish(kontrakt, baza):
    """Mahsulot aytilmagan — bu xato emas, aniqlashtirish."""
    llm = temur_javobi(qatorlar=[], sorash_kerak=["Qaysi mahsulot kerak?"])
    k = await temur_yasa(kontrakt, baza, llm).ishla("KP tayyorla, ruscha")

    assert k.holat is Holat.ANIQLIK_KERAK
    assert k.natija["savollar"] == ["Qaysi mahsulot kerak?"]


@pytest.mark.asyncio
async def test_katalog_ishlamasa_ham_kp_tuziladi(kontrakt, baza, monkeypatch, tmp_path):
    """Katalog yiqilsa foydalanuvchi bloklanmaydi — KP baribir chiqadi."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    llm = temur_javobi(
        qatorlar=[{
            "nomi": "Гибкая вставка", "spetsifikatsiya": "", "miqdor": 5,
            "birlik": "dona", "katalogda_yoq": True, "izoh": "",
        }],
        topilmagan_mahsulotlar=["Гибкая вставка"],
    )
    # `soxta_api()` — hech narsa sozlanmagan, ya'ni ApiXatosi (katalog yiqilgan)
    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=llm, baza=baza,
        api=soxta_api(), qidiruv_manbasi=soxta_bilim(("KP shabloni", "matn")),
    )
    k = await temur.ishla("5 dona Гибкая вставка uchun KP, ruscha")

    # KP tuzildi, foydalanuvchi bloklanmadi
    assert k.holat is Holat.TUGADI
    assert k.natija["raqam"]
    assert len(k.natija["qatorlar"]) == 1
    # Lekin tasdiqlanmagani ochiq aytilgan
    assert any("tasdiqlanmadi" in o for o in k.natija["ogohlantirishlar"])

    # Modelga "baribir tuz" deb aytilgan
    xabar = llm.chaqiruvlar[0]["messages"][0]["content"]
    assert "BARIBIR tuz" in xabar


@pytest.mark.asyncio
async def test_promptda_chegaralar():
    from app.agentlar.tijorat_menejeri import TIZIM_PROMPT

    assert "NARX YOZMAYSAN" in TIZIM_PROMPT
    assert "SPESIFIKATSIYANI O'YLAB TOPMAYSAN" in TIZIM_PROMPT
    assert "YUBORMAYSAN" in TIZIM_PROMPT
    assert "TAXMIN QILMAYSAN" in TIZIM_PROMPT


# --- ko'rinish ---------------------------------------------------------------


def test_kp_odamcha_korinadi():
    natija = {
        "raqam": "KP-2026-0042",
        "sana": "2026-07-29",
        "mijoz": {"nomi": "Oqtepa MChJ", "aloqa": "+998901112233"},
        "qatorlar": [
            {"nomi": "ВК-250П", "spetsifikatsiya": "230 V", "miqdor": 5,
             "birlik": "dona", "birlik_narx": 2_400_000, "jami": 12_000_000},
            {"nomi": "КПР-1", "spetsifikatsiya": "", "miqdor": 2,
             "birlik": "dona", "birlik_narx": None, "jami": None},
        ],
        "summa": 12_000_000, "qqs": 1_440_000, "jami": 13_440_000, "valyuta": "UZS",
        "ogohlantirishlar": ["1 ta mahsulot narxi topilmadi"],
        "fayllar": {"docx": "x.docx", "pdf": "x.pdf"},
        "sorash_kerak": [],
        "taklif_qilingan": [],
    }
    matn = agent_matni("proposal-builder", natija)

    assert "Tijorat taklifi № KP-2026-0042" in matn
    assert "5 dona × 2 400 000 UZS = 12 000 000 UZS" in matn
    assert "[narx to'ldirilishi kerak]" in matn
    assert "Umumiy summa: 13 440 000 UZS" in matn
    assert "DOCX, PDF" in matn
    assert "Mijozga yuborilmadi" in matn
    # Texnik maydonlar yo'q
    for belgi in ("{", "}", '"nomi"', "None"):
        assert belgi not in matn


# --- KP tili so'raladi -------------------------------------------------------


@pytest.mark.asyncio
async def test_til_aytilmagan_bolsa_soraydi(kontrakt, baza, monkeypatch, tmp_path):
    """Hujjat mijozga ketadi — tilni taxmin qilmaymiz, so'raymiz.

    Avval jimgina rus tili olinardi. Menejer o'zbekcha xohlasa, buni
    faqat tayyor KP ni ochib ko'rgandan keyin bilardi.
    """
    from app.agentlar.tijorat_menejeri import TIL_SAVOLI

    menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "Aziz Karimov", "+998 90 111 11 11")

    k = await temur_yasa(kontrakt, baza, temur_javobi()).ishla(
        "5 dona ВК-250П uchun KP", {"telegram_id": 111}
    )

    assert k.holat is Holat.ANIQLIK_KERAK
    assert k.natija["savollar"] == [TIL_SAVOLI]
    # Bot javobni qaysi savolga bog'lashini shu belgidan biladi
    assert k.natija["kerak"] == "til"
    assert "ruscha" in TIL_SAVOLI and "o'zbekcha" in TIL_SAVOLI


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("sorov", "kutilgan"),
    [
        ("KP ni ruscha qilib ber", "ru"),
        ("KP ni o'zbekcha qilib ber", "uz"),
        ("KP на русском", "ru"),
    ],
)
async def test_til_aytilgan_bolsa_soralmaydi(
    kontrakt, baza, monkeypatch, tmp_path, sorov, kutilgan
):
    menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "Aziz Karimov", "+998 90 111 11 11")

    llm = temur_javobi(til=kutilgan)
    k = await temur_yasa(kontrakt, baza, llm).ishla(sorov, {"telegram_id": 111})

    assert k.holat is not Holat.ANIQLIK_KERAK
    assert f"HUJJAT TILI: {kutilgan}" in llm.chaqiruvlar[0]["messages"][0]["content"]


@pytest.mark.asyncio
async def test_ozbekcha_kp_da_ism_lotinda(kontrakt, baza, monkeypatch, tmp_path):
    """Kirillda kiritilgan ism o'zbekcha KP da lotinga o'giriladi."""
    yozilgan = menejer_sozla(monkeypatch, tmp_path)
    await baza.menejer_yoz(111, "\u0410\u0437\u0438\u0437 \u041a\u0430\u0440\u0438\u043c\u043e\u0432", "+998 90 111 11 11")

    await temur_yasa(kontrakt, baza, temur_javobi()).ishla(
        "KP tayyorla, o'zbekcha", {"telegram_id": 111}
    )

    assert yozilgan["kp"].til == "uz"
    assert yozilgan["kp"].rekvizitlar["menejer"] == "Aziz Karimov"


@pytest.mark.asyncio
async def test_qolda_yozilgan_variant_kp_da_ustun(
    kontrakt, baza, monkeypatch, tmp_path
):
    """Pasportdagi yozuvni avtomatik o'girish bosib ketmasin.

    `rekvizitlar.yaml` da `ism_kiril` bo'lsa, KP da aynan shu turadi.
    """
    import app.agentlar.tijorat_menejeri as modul

    yozilgan = menejer_sozla(monkeypatch, tmp_path)
    xodim = {
        "ism": "Shukhrat Yuldashev",
        "telefon": "+998 90 444 44 44",
        "ism_kiril": "\u0428\u0443\u0445\u0440\u0430\u0442 \u042e\u043b\u0434\u0430\u0448\u0435\u0432",
    }
    monkeypatch.setattr(modul, "menejer_uchun", lambda tg_id: xodim)

    await temur_yasa(kontrakt, baza, temur_javobi()).ishla(
        "KP ruscha", {"telegram_id": 111}
    )

    # Avtomatik o'girish "\u0428\u0443\u043a\u04b3\u0440\u0430\u0442" berardi \u2014 qo'lda yozilgani ustun.
    assert yozilgan["kp"].rekvizitlar["menejer"] == "\u0428\u0443\u0445\u0440\u0430\u0442 \u042e\u043b\u0434\u0430\u0448\u0435\u0432"


# --- dollardagi variant narxi: kurs KP da ochiq yozilsin --------------------


def _docx_matni(yol: str) -> str:
    """Tayyor hujjatning butun matni — mijoz aynan shuni ko'radi."""
    from docx import Document

    hujjat = Document(yol)
    qismlar = [p.text for p in hujjat.paragraphs]
    for jadval in hujjat.tables:
        for qator in jadval.rows:
            qismlar += [k.text for k in qator.cells]
    return chr(10).join(qismlar)


@pytest.mark.asyncio
async def test_dollardan_hisoblangan_narxda_kurs_KP_da_yoziladi(
    kontrakt, baza, monkeypatch, tmp_path
):
    """Mijoz raqam qayerdan chiqqanini ko'rsin — kurs o'zgarganda bahs bo'lmasin."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=temur_javobi(), baza=baza,
        api=soxta_api(qidir_keng=KATALOG,
                      model_narxi=(2_391_960, "ВК-250П (199.33 $ x 12000)")),
        qidiruv_manbasi=soxta_bilim(("KP shabloni", "matn")),
    )
    k = await temur.ishla("5 dona ВК-250П uchun KP, ruscha")

    matn = _docx_matni(k.natija["fayllar"]["docx"])
    # KP ruscha — izoh ham ruscha chiqadi. Kurs raqami ikkala tilda ham bir xil.
    assert "12 000" in matn and ("курсу" in matn or "kursida" in matn), matn[-500:]


@pytest.mark.asyncio
async def test_somdagi_narxda_kurs_yozilmaydi(kontrakt, baza, monkeypatch, tmp_path):
    """Eski `models[].price` SO'MDA — unga kurs aloqador emas."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=temur_javobi(), baza=baza,
        api=soxta_api(qidir_keng=KATALOG, model_narxi=(999_000, "ВК-250П")),
        qidiruv_manbasi=soxta_bilim(("KP shabloni", "matn")),
    )
    k = await temur.ishla("5 dona ВК-250П uchun KP, ruscha")

    matn = _docx_matni(k.natija["fayllar"]["docx"])
    assert "курсу" not in matn and "kursida" not in matn


@pytest.mark.asyncio
async def test_quvvat_aytilmasa_menejer_OGOHLANTIRILADI(
    kontrakt, baza, monkeypatch, tmp_path
):
    """Eng arzon variant jim tanlanmasin — 156 $ va 199 $ farqi katta."""
    import app.agentlar.tijorat_menejeri as modul

    monkeypatch.setattr(modul, "narxlar", narx_royxati)
    monkeypatch.setattr(modul, "rekvizitlar", lambda: {"nomi": "Test MChJ"})
    monkeypatch.setattr(modul.sozlama(), "kp_yoli", str(tmp_path), raising=False)

    temur = TijoratMenejeri(
        kontrakt=kontrakt, llm=temur_javobi(), baza=baza,
        api=soxta_api(qidir_keng=KATALOG, model_narxi=(
            1_881_120, "ВК-250П-0,12/1500 (156.76 $ x 12000) — 5 variantdan eng arzoni")),
        qidiruv_manbasi=soxta_bilim(("KP shabloni", "matn")),
    )
    k = await temur.ishla("5 dona ВК-250П uchun KP, ruscha")

    assert any("eng arzoni" in o for o in k.natija["ogohlantirishlar"]), \
        k.natija["ogohlantirishlar"]
